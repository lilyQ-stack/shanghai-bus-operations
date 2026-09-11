import argparse
import csv
import json
import math
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
TZ = timezone(timedelta(hours=8))
ROUTES = ["浦东78路", "浦东35路", "182路"]
EARLY_REVERSE_RATIO = 0.80
MIN_REVERSE_ELAPSED_MIN = 20

FIELDS = [
    "线路", "车牌号", "发车时间", "发车站", "终点站", "班次类型", "区间/异常说明",
    "预计到达时间", "全程时间（分钟）", "到达置信度", "到达估算方法", "参与车速排名",
    "最后可靠采集站点", "最后可靠预计到达时间", "本车首次观测",
]


def parse_hm(text):
    if not text:
        return None
    m = re.fullmatch(r"(\d{1,2}):(\d{2})", str(text).strip())
    if not m:
        return None
    h, minute = map(int, m.groups())
    if not (0 <= h < 24 and 0 <= minute < 60):
        return None
    return h * 60 + minute


def fmt_hm(value):
    if value is None:
        return ""
    value %= 24 * 60
    return f"{value // 60:02d}:{value % 60:02d}"


def minutes_between(start, end):
    if start is None or end is None:
        return None
    diff = end - start
    if diff < 0:
        diff += 24 * 60
    return diff


def median(values):
    values = sorted(v for v in values if v is not None)
    if not values:
        return None
    n = len(values)
    if n % 2:
        return values[n // 2]
    return (values[n // 2 - 1] + values[n // 2]) / 2


def load_snapshots(date):
    snapshots, source_files = [], []
    for route in ROUTES:
        path = ROOT / "data" / "shmaas" / f"{date}-{route}.jsonl"
        if not path.exists():
            continue
        source_files.append(str(path.relative_to(ROOT)))
        with path.open(encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    rec = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if rec.get("success") and rec.get("date_cst") == date:
                    snapshots.append(rec)
    return snapshots, source_files


def event_time(snapshot):
    raw = snapshot.get("sample_time_cst")
    try:
        return datetime.fromisoformat(raw).astimezone(TZ)
    except Exception:
        return None


def collect_evidence(snapshots):
    route_info = defaultdict(dict)
    events = defaultdict(list)
    dispatches = defaultdict(set)
    first_seen = {}

    for snap in snapshots:
        route = snap.get("route")
        dt = event_time(snap)
        if not route or not dt:
            continue
        sample_min = dt.hour * 60 + dt.minute + dt.second / 60

        for direction in snap.get("directions", []):
            d = direction.get("direction")
            if d is None:
                continue
            route_info[(route, d)] = {
                "start": direction.get("start_stop", ""),
                "end": direction.get("end_stop", ""),
                "stop_count": direction.get("stop_count") or 0,
            }

        for vehicle in snap.get("vehicles", []):
            plate = vehicle.get("plate")
            if not plate:
                continue
            first_seen[(route, plate)] = min(first_seen.get((route, plate), dt), dt)
            for obs in vehicle.get("observations", []):
                d = obs.get("direction")
                if d is None:
                    continue
                role = obs.get("role") or ""
                dispatch = parse_hm(obs.get("dispatch_time"))
                if role == "dispatch" and dispatch is not None:
                    dispatches[(route, plate, d)].add(dispatch)
                    # A dispatch observation is a timetable/dispatch plan, not proof that the
                    # vehicle is physically travelling in this direction at sample_time.
                    # Keep it for trip boundaries only; never mix it into trajectory evidence.
                    continue
                eta = obs.get("eta_min")
                try:
                    eta = float(eta) if eta is not None else None
                except Exception:
                    eta = None
                distance = obs.get("distance_m")
                try:
                    distance = float(distance) if distance is not None else None
                except Exception:
                    distance = None
                seq = obs.get("stop_seq")
                try:
                    seq = int(seq) if seq is not None else None
                except Exception:
                    seq = None
                events[(route, plate, d)].append({
                    "time": sample_min,
                    "dt": dt,
                    "role": role,
                    "stop_seq": seq,
                    "stop_name": obs.get("stop_name", ""),
                    "eta": eta,
                    "distance": distance,
                    "service_hints": obs.get("service_hints") or [],
                })

    for key in events:
        events[key].sort(key=lambda x: x["time"])
    return route_info, events, dispatches, first_seen


def explicit_service_hint(trip_events):
    for ev in trip_events:
        hints = [str(x) for x in ev.get("service_hints") or []]
        joined = " ".join(hints)
        if any(token in joined for token in ("区间", "短线", "折返")):
            return "疑似区间车", joined
    return None, ""


def derive_min_full_runtimes(route_info, events, dispatches):
    candidates = defaultdict(list)
    for (route, plate, d), deps in dispatches.items():
        info = route_info.get((route, d), {})
        stop_count = info.get("stop_count") or 0
        terminal = info.get("end") or ""
        same = events.get((route, plate, d), [])
        dep_list = sorted(deps)
        for i, dep in enumerate(dep_list):
            upper = dep_list[i + 1] if i + 1 < len(dep_list) else dep + 240
            trip = [e for e in same if dep <= e["time"] < upper]
            for ev in trip:
                near_terminal = (
                    ev.get("stop_name") == terminal
                    or (stop_count and ev.get("stop_seq") and ev["stop_seq"] / stop_count >= 0.85)
                )
                if not near_terminal:
                    continue
                eta = ev.get("eta")
                predicted = ev["time"] + eta if eta is not None else ev["time"]
                runtime = predicted - dep
                if 30 <= runtime <= 240:
                    candidates[(route, d)].append(runtime)
    return {key: min(vals) for key, vals in candidates.items() if vals}


def terminal_eta(trip_events, terminal, stop_count):
    candidates = []
    confirmed_candidates = []
    for ev in trip_events:
        seq = ev.get("stop_seq")
        near_terminal = ev.get("stop_name") == terminal or (
            stop_count and seq and seq / stop_count >= 0.85
        )
        if not near_terminal:
            continue
        eta = ev.get("eta")
        if eta is None:
            continue
        pred = ev["time"] + eta
        candidates.append(pred)
        confirmed = (
            ev.get("stop_name") == terminal
            and (eta <= 1 or (ev.get("distance") is not None and ev["distance"] <= 100))
        )
        if confirmed:
            confirmed_candidates.append(pred)

    if confirmed_candidates:
        p = median(confirmed_candidates)
        spread = max(confirmed_candidates) - min(confirmed_candidates) if len(confirmed_candidates) > 1 else 2
        return p, "终点近站ETA确认", max(2, spread / 2), True
    if len(candidates) >= 2:
        p = median(candidates)
        spread = max(candidates) - min(candidates)
        return p, "终点ETA多样本共识", max(5, spread / 2), False
    if candidates:
        return candidates[-1], "末次终点ETA", 12, False
    return None, "", None, False


def confidence_grade(error):
    if error is None:
        return "D"
    if error <= 5:
        return "A"
    if error <= 10:
        return "B"
    if error <= 15:
        return "C"
    return "D"


def trajectory_classification(route, plate, d, dep, next_dep, route_info, events, min_full):
    info = route_info.get((route, d), {})
    stop_count = info.get("stop_count") or 0
    terminal = info.get("end") or ""
    same = [e for e in events.get((route, plate, d), []) if dep <= e["time"] < next_dep]
    opposite = [
        e for e in events.get((route, plate, 1 - d), [])
        if dep + MIN_REVERSE_ELAPSED_MIN <= e["time"] < next_dep
    ]
    if not same:
        return "运行异常待查", "发车后缺少同向轨迹", "", ""

    hint_type, hint_text = explicit_service_hint(same)
    if hint_type:
        last = same[-1]
        return hint_type, f"实时服务提示：{hint_text}", last["stop_name"], ""

    first_reverse = opposite[0] if opposite else None
    last_same = max((e for e in same if not first_reverse or e["time"] <= first_reverse["time"]), key=lambda e: e["time"], default=same[-1])
    seq = last_same.get("stop_seq") or 0
    near_terminal = last_same.get("stop_name") == terminal or (stop_count and seq / stop_count >= 0.85)

    if first_reverse:
        elapsed = first_reverse["time"] - dep
        baseline = min_full.get((route, d))
        if baseline is not None and elapsed < baseline * EARLY_REVERSE_RATIO:
            if not near_terminal:
                return (
                    "疑似区间车",
                    f"同车{elapsed:.0f}分钟后已反向运行，早于同方向最短全程{baseline:.0f}分钟的80%阈值",
                    last_same["stop_name"], ""
                )
            return (
                "运行异常待查",
                f"同车较早切换方向，但末次同向观测已进入线路末段，可能为终点漏采",
                last_same["stop_name"], ""
            )
        if baseline is None and not near_terminal and first_reverse["time"] - last_same["time"] <= 20:
            return (
                "疑似区间车",
                "缺少当日全程基线；同车在中途末次观测后20分钟内出现反向运行",
                last_same["stop_name"], ""
            )
        if near_terminal:
            return (
                "全程车",
                "同车随后反向运行，且末次同向观测已进入线路末段；按全程运行处理",
                last_same["stop_name"], ""
            )
        return "运行异常待查", "同车后续方向切换，但未满足提前折返判据", last_same["stop_name"], ""

    if near_terminal:
        predicted = last_same["time"] + last_same["eta"] if last_same.get("eta") is not None else None
        return (
            "全程车",
            "末次同向观测已进入线路末段，且未满足提前反向区间车判据；按全程运行处理",
            last_same["stop_name"], fmt_hm(predicted),
        )
    return "运行中待确认", "尚未观察到终点或反向运行证据", last_same["stop_name"], ""


def build_rows(date, route_info, events, dispatches, first_seen):
    min_full = derive_min_full_runtimes(route_info, events, dispatches)
    rows = []
    for (route, plate, d), deps in sorted(dispatches.items()):
        info = route_info.get((route, d), {})
        terminal = info.get("end") or ""
        stop_count = info.get("stop_count") or 0
        dep_list = sorted(deps)
        for i, dep in enumerate(dep_list):
            next_dep = dep_list[i + 1] if i + 1 < len(dep_list) else dep + 240
            trip_events = [e for e in events.get((route, plate, d), []) if dep <= e["time"] < next_dep]
            service_type, reason, last_stop, last_eta = trajectory_classification(
                route, plate, d, dep, next_dep, route_info, events, min_full
            )
            arrival = None
            method = ""
            error = None
            confirmed = False
            if service_type in ("全程车", "运行中待确认"):
                arrival, method, error, confirmed = terminal_eta(trip_events, terminal, stop_count)

            if arrival is not None and confirmed and next_dep < dep + 240 and next_dep <= arrival:
                arrival = next_dep - 2
                method = "终点确认到达后按下班反向/同向发车边界校正"
                error = max(error or 5, 5)

            runtime = minutes_between(dep, arrival) if arrival is not None else None
            if runtime is not None and (runtime <= 0 or runtime > 240):
                arrival = None
                runtime = None
                method = ""
                error = None

            grade = confidence_grade(error)
            eligible = service_type == "全程车" and runtime is not None and grade in ("A", "B", "C")
            if trip_events:
                last = trip_events[-1]
                if not last_stop:
                    last_stop = last.get("stop_name", "")
                if not last_eta and last.get("eta") is not None:
                    last_eta = fmt_hm(last["time"] + last["eta"])
            rows.append({
                "线路": route,
                "车牌号": plate,
                "发车时间": fmt_hm(dep),
                "发车站": info.get("start", ""),
                "终点站": terminal,
                "班次类型": service_type,
                "区间/异常说明": reason,
                "预计到达时间": fmt_hm(arrival) if arrival is not None else "",
                "全程时间（分钟）": round(runtime, 1) if runtime is not None else "",
                "到达置信度": grade if arrival is not None else "D",
                "到达估算方法": method,
                "参与车速排名": "是" if eligible else "否",
                "最后可靠采集站点": last_stop,
                "最后可靠预计到达时间": last_eta,
                "本车首次观测": first_seen.get((route, plate)).strftime("%H:%M") if first_seen.get((route, plate)) else "",
                "_终点确认到达": confirmed,
            })
    return rows


def rank_rows(rows):
    runtimes = defaultdict(list)
    for row in rows:
        if row["参与车速排名"] == "是" and row["全程时间（分钟）"] != "":
            runtimes[(row["线路"], row["车牌号"])].append(float(row["全程时间（分钟）"]))
    averages = {key: sum(vals) / len(vals) for key, vals in runtimes.items() if vals}
    ranks = {}
    for route in ROUTES:
        ordered = sorted((avg, plate) for (r, plate), avg in averages.items() if r == route)
        for rank, (avg, plate) in enumerate(ordered, 1):
            ranks[(route, plate)] = rank
    for row in rows:
        row["当日车速排名"] = ranks.get((row["线路"], row["车牌号"]), "")
    return averages


def write_csv(date, rows):
    out_dir = ROOT / "data" / "export"
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_fields = FIELDS + ["当日车速排名"]
    for route in ROUTES:
        path = out_dir / f"{date}-{route}.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=csv_fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows([r for r in rows if r["线路"] == route])


def write_xlsx(date, rows, averages):
    out_dir = ROOT / "data" / "export"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{date}-上海公交运营分析.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["线路", "车辆数", "班次数", "可排名班次数", "最快车辆", "平均全程时间（分钟）"])
    for route in ROUTES:
        rr = [r for r in rows if r["线路"] == route]
        plates = sorted({r["车牌号"] for r in rr})
        eligible = [r for r in rr if r["参与车速排名"] == "是"]
        ranked = sorted((avg, plate) for (r, plate), avg in averages.items() if r == route)
        fastest = ranked[0][1] if ranked else ""
        avg = ranked[0][0] if ranked else ""
        ws.append([route, len(plates), len(rr), len(eligible), fastest, round(avg, 1) if avg != "" else ""])

    headers = FIELDS + ["当日车速排名"]
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for sheet_name, route_filter in [("全部班次", None)] + [(r, r) for r in ROUTES]:
        sh = wb.create_sheet(sheet_name)
        sh.append(headers)
        subset = rows if route_filter is None else [r for r in rows if r["线路"] == route_filter]
        for row in subset:
            sh.append([row.get(h, "") for h in headers])
            if row["班次类型"] != "全程车":
                for cell in sh[sh.max_row]:
                    cell.fill = yellow
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for col in sh.columns:
            letter = col[0].column_letter
            width = min(max(len(str(c.value or "")) for c in col) + 2, 42)
            sh.column_dimensions[letter].width = max(width, 10)
    wb.save(path)
    return path


def write_meta(date, snapshots, source_files, rows, averages):
    out_dir = ROOT / "data" / "export"
    path = out_dir / f"{date}-analysis-meta.json"
    payload = {
        "date": date,
        "generated_at_cst": datetime.now(TZ).isoformat(timespec="seconds"),
        "source": "SHMAAS anonymous public endpoints only",
        "source_files": source_files,
        "successful_snapshots": len(snapshots),
        "rules": {
            "short_turn": (
                "A trip is flagged only when the same plate is observed in actual current/next trajectory evidence "
                "in the reverse direction after at least 20 minutes and before the next same-direction planned departure. "
                "Planned dispatch observations are excluded from trajectory/reverse evidence. "
                "If a same-day minimum full-trip runtime exists, reverse must occur before 80% of it and before the vehicle reaches the route end segment. "
                "Without a baseline, reverse must follow an intermediate last same-direction observation within 20 minutes."
            ),
            "arrival_rule": (
                "Terminal ETA evidence is preferred. Only a near-terminal observation with ETA <=1 minute or distance <=100m "
                "is treated as confirmed enough to override an impossible next planned departure. A terminal sampling gap by itself "
                "does not make the service abnormal; service classification and arrival confidence are evaluated separately. "
                "Confidence A/B/C/D corresponds to estimated uncertainty <=5, <=10, <=15, >15 or insufficient minutes."
            ),
            "ranking": (
                "Within each route, arithmetic mean full-trip runtime per vehicle over eligible full trips; "
                "lower runtime ranks faster. Short-turn/abnormal trips and D-confidence arrivals are excluded."
            ),
        },
        "eligible_vehicle_average_runtime": {
            f"{route}|{plate}": round(avg, 2) for (route, plate), avg in sorted(averages.items())
        },
        "row_count": len(rows),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def main():
    parser = argparse.ArgumentParser(description="Build daily SHMAAS-only bus operations exports.")
    parser.add_argument("--date", help="Shanghai calendar date to analyze (YYYY-MM-DD). Defaults to today in Asia/Shanghai.")
    args = parser.parse_args()
    date = args.date or datetime.now(TZ).date().isoformat()
    try:
        datetime.strptime(date, "%Y-%m-%d")
    except ValueError as exc:
        parser.error(f"invalid --date {date!r}; expected YYYY-MM-DD")
    snapshots, source_files = load_snapshots(date)
    if not snapshots:
        raise SystemExit(f"No successful SHMAAS snapshots found for {date}")
    route_info, events, dispatches, first_seen = collect_evidence(snapshots)
    rows = build_rows(date, route_info, events, dispatches, first_seen)
    averages = rank_rows(rows)
    write_csv(date, rows)
    xlsx = write_xlsx(date, rows, averages)
    meta = write_meta(date, snapshots, source_files, rows, averages)
    print(json.dumps({
        "date": date,
        "snapshots": len(snapshots),
        "rows": len(rows),
        "xlsx": str(xlsx.relative_to(ROOT)),
        "meta": str(meta.relative_to(ROOT)),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
