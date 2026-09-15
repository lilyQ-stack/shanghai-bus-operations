from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TZ = ZoneInfo("Asia/Shanghai")
YELLOW = PatternFill(fill_type="solid", fgColor="FFF2CC")
HEADER = PatternFill(fill_type="solid", fgColor="D9EAF7")
HIDDEN_EXPORT_FIELDS = {"方向"}


def read_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_xlsx(csv_path: Path):
    rows = read_csv(csv_path)
    if not rows:
        return None
    fields = [f for f in rows[0].keys() if f not in HIDDEN_EXPORT_FIELDS]
    wb = Workbook()
    ws = wb.active
    ws.title = "运营数据"
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ws.append([row.get(f, "") for f in fields])
        excel_row = ws.max_row
        if row.get("班次类型") == "运行异常待查":
            for field in ("最后可靠采集站点", "最后可靠预计到达时间"):
                if field in fields:
                    cell = ws.cell(excel_row, fields.index(field) + 1)
                    cell.fill = YELLOW
                    cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    for idx, field in enumerate(fields, 1):
        width = min(36, max(10, len(field) * 2 + 2))
        for cell in ws[get_column_letter(idx)][1:]:
            if cell.value is not None:
                width = min(36, max(width, len(str(cell.value)) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    out = csv_path.with_suffix(".xlsx")
    wb.save(out)
    print(out)
    return out


def main():
    date = datetime.now(TZ).date().isoformat()
    export = ROOT / "data" / "export"
    for path in sorted(export.glob(f"{date}-*.csv")):
        if path.name.endswith("-vehicles.csv") or path.name.endswith("-vehicle-timeline.csv"):
            continue
        write_xlsx(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
